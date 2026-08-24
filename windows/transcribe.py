#!/usr/bin/env python3
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document


FILLER_PATTERNS = [
    r"\b(呃+|嗯+|啊+|额+|这个|那个|然后然后|就是就是)\b",
    r"(哈哈哈+|呵呵+|嘿嘿+|笑声|掌声|音乐)",
    r"\b(um+|uh+|erm+|like you know)\b",
]


def emit(message):
    print(message, file=sys.stderr, flush=True)


def find_ffmpeg():
    candidates = [
        os.environ.get("TRANSCRIBE_STUDIO_FFMPEG"),
        "/opt/homebrew/bin/ffmpeg",
        "/usr/local/bin/ffmpeg",
        shutil.which("ffmpeg"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    raise RuntimeError("未找到 ffmpeg。请先安装 Homebrew ffmpeg，或设置 TRANSCRIBE_STUDIO_FFMPEG。")


def find_ffprobe(ffmpeg):
    env_path = os.environ.get("TRANSCRIBE_STUDIO_FFPROBE")
    candidates = [
        env_path,
        str(Path(ffmpeg).with_name("ffprobe")) if ffmpeg else None,
        "/opt/homebrew/bin/ffprobe",
        "/usr/local/bin/ffprobe",
        shutil.which("ffprobe"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    raise RuntimeError("未找到 ffprobe。请确认 ffmpeg 安装完整。")


def extract_audio(ffmpeg, input_path, output_path):
    emit("正在抽取音频...")
    probe_cmd = [
        find_ffprobe(ffmpeg),
        "-v",
        "error",
        "-select_streams",
        "a",
        "-show_entries",
        "stream=index",
        "-of",
        "csv=p=0",
        str(input_path),
    ]
    probe = subprocess.run(probe_cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if not probe.stdout.strip():
        raise RuntimeError(
            "录制文件没有音频轨，无法转写。请在 macOS 的“屏幕与系统音频录制”和“麦克风”中授权本应用；"
            "如果仍无系统声音，请打开电脑外放，让应用使用麦克风兜底录音。"
        )
    cmd = [
        ffmpeg,
        "-y",
        "-i",
        str(input_path),
        "-vn",
        "-ac",
        "1",
        "-ar",
        "16000",
        "-f",
        "wav",
        str(output_path),
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)


def load_engine(model_name, language):
    try:
        from faster_whisper import WhisperModel

        def run(audio_path):
            emit(f"使用 faster-whisper 模型：{model_name}")
            model = WhisperModel(model_name, device="auto", compute_type="auto")
            kwargs = {"vad_filter": True}
            if language and language != "auto":
                kwargs["language"] = language
            segments, info = model.transcribe(str(audio_path), **kwargs)
            return [
                {
                    "start": float(segment.start),
                    "end": float(segment.end),
                    "text": segment.text.strip(),
                }
                for segment in segments
            ], getattr(info, "language", language or "auto")

        return run
    except Exception as first_error:
        try:
            import whisper

            def run(audio_path):
                emit(f"使用 openai-whisper 模型：{model_name}")
                model = whisper.load_model(model_name)
                kwargs = {}
                if language and language != "auto":
                    kwargs["language"] = language
                result = model.transcribe(str(audio_path), **kwargs)
                return [
                    {
                        "start": float(segment.get("start", 0)),
                        "end": float(segment.get("end", 0)),
                        "text": segment.get("text", "").strip(),
                    }
                    for segment in result.get("segments", [])
                ], result.get("language", language or "auto")

            return run
        except Exception as second_error:
            raise RuntimeError(
                "未检测到本机 Whisper AI 转写引擎。可执行：\n"
                "python3 -m pip install faster-whisper python-docx\n"
                "或：python3 -m pip install openai-whisper python-docx\n"
                f"faster-whisper 错误：{first_error}\nopenai-whisper 错误：{second_error}"
            )


def normalize_text(text):
    text = re.sub(r"\s+", " ", text).strip()
    for pattern in FILLER_PATTERNS:
        text = re.sub(pattern, "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"([。！？!?])\1+", r"\1", text)
    return text


def clean_segments(segments, style):
    cleaned = []
    for segment in segments:
        text = normalize_text(segment["text"])
        if not text or len(text) <= 1:
            continue
        if style != "verbatim" and re.fullmatch(r"[啊嗯呃额哈\s，。,.!?！？]+", text):
            continue
        cleaned.append({**segment, "text": text})
    return cleaned


def split_sentences(text):
    parts = re.split(r"(?<=[。！？!?])\s*", text)
    return [part.strip() for part in parts if part.strip()]


def build_document(cleaned, style):
    full_text = "\n".join(segment["text"] for segment in cleaned).strip()
    sentences = split_sentences(full_text)
    if style == "outline":
        important = sentences[:80]
        body = "\n".join(f"- {sentence}" for sentence in important)
    elif style == "verbatim":
        body = full_text
    else:
        important = []
        for sentence in sentences:
            if len(sentence) < 6:
                continue
            if sentence not in important:
                important.append(sentence)
        body = "\n".join(important)
    return body.strip() or full_text


def write_outputs(output_dir, source_path, cleaned, language, style):
    title = source_path.stem
    body = build_document(cleaned, style)
    timeline = "\n".join(
        f"[{int(item['start'] // 60):02d}:{int(item['start'] % 60):02d}] {item['text']}"
        for item in cleaned
    )

    markdown = output_dir / "转写文档.md"
    txt = output_dir / "转写文档.txt"
    docx = output_dir / "转写文档.docx"

    md_content = (
        f"# {title}\n\n"
        f"## 整理稿\n\n{body}\n\n"
        f"## 时间线原文\n\n{timeline}\n"
    )
    markdown.write_text(md_content, encoding="utf-8")
    txt.write_text(f"{title}\n\n整理稿\n\n{body}\n\n时间线原文\n\n{timeline}\n", encoding="utf-8")

    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_heading("整理稿", level=2)
    for paragraph in body.splitlines():
        if paragraph.startswith("- "):
            doc.add_paragraph(paragraph[2:], style="List Bullet")
        elif paragraph.strip():
            doc.add_paragraph(paragraph)
    doc.add_heading("时间线原文", level=2)
    for line in timeline.splitlines():
        doc.add_paragraph(line)
    doc.add_paragraph(f"识别语言：{language}；整理方式：{style}")
    doc.save(docx)

    return markdown, txt, docx, body


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--language", default="zh")
    parser.add_argument("--model", default="small")
    parser.add_argument("--style", default="clean")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    ffmpeg = find_ffmpeg()
    audio_path = output_dir / "audio.wav"
    extract_audio(ffmpeg, input_path, audio_path)

    engine = load_engine(args.model, args.language)
    emit("正在进行 AI 语音识别...")
    segments, detected_language = engine(audio_path)
    cleaned = clean_segments(segments, args.style)

    raw_json = output_dir / "segments.json"
    raw_json.write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")

    markdown, txt, docx, body = write_outputs(output_dir, input_path, cleaned, detected_language, args.style)
    summary = body.splitlines()[0][:120] if body else "已完成转写。"
    print(json.dumps({
        "output_dir": str(output_dir),
        "markdown": str(markdown),
        "txt": str(txt),
        "docx": str(docx),
        "segments": str(raw_json),
        "summary": summary,
        "language": detected_language,
    }, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    try:
        main()
    except subprocess.CalledProcessError as error:
        raise SystemExit(f"音频处理失败：{error.stderr.decode('utf-8', errors='ignore')}")
    except Exception as error:
        raise SystemExit(str(error))
