#!/usr/bin/env python3
import argparse
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import textwrap
import urllib.error
import urllib.request
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document


FILLER_PATTERNS = [
    r"\b(呃+|嗯+|啊+|额+|这个|那个|然后然后|就是就是)\b",
    r"(哈哈哈+|呵呵+|嘿嘿+|笑声|掌声|音乐)",
    r"\b(um+|uh+|erm+|like you know)\b",
]

FALLBACK_T2S = str.maketrans({
    "開": "开", "發": "发", "這": "这", "個": "个", "結": "结", "果": "果", "沒": "没",
    "為": "为", "實": "实", "驗": "验", "證": "证", "關": "关", "鍵": "键",
    "題": "题", "願": "愿", "壯": "壮", "談": "谈", "論": "论", "據": "据",
    "瘋": "疯", "點": "点", "應": "应", "講": "讲", "觀": "观", "別": "别",
    "線": "线", "買": "买", "殘": "残", "酷": "酷", "該": "该", "訪": "访",
})

_OPENCC_CONVERTER = None


SUBTITLE_REGION_FILTER = "crop=iw:ih*0.32:0:ih*0.62,scale=1280:-1,format=gray"
PROGRESS_PREFIX = "__LC_PROGRESS__"

DOCUMENT_TEMPLATES = {
    "general": {
        "label": "通用整理",
        "hint": "输出一份适合阅读和存档的结构化整理稿，包含主题概述、重点内容、结论和行动建议。"
    },
    "live_recap": {
        "label": "直播复盘",
        "hint": "按“核心观点、精彩案例、观众提问、成交话术、待优化项、可复用素材”组织内容。适合知识付费讲师直播后复盘。"
    },
    "course_notes": {
        "label": "课程笔记",
        "hint": "按“知识图谱、核心概念、方法论步骤、关键案例、易错点、课后行动”组织内容。"
    },
    "material_extract": {
        "label": "素材提取",
        "hint": "按“可切片片段、爆款开头、情绪高光、金句、短视频标题、剪辑建议”组织内容。"
    },
    "sales_script": {
        "label": "成交话术",
        "hint": "按“用户痛点、信任建立、价值呈现、异议处理、成交话术、可复用表达”组织内容。"
    },
    "quotes": {
        "label": "金句提取",
        "hint": "提取有传播价值的金句，并补充适用场景和原文上下文。不要为了凑数量编造金句。"
    },
}


def emit(message):
    print(message, file=sys.stderr, flush=True)


def emit_progress(stage, percent, message="", current=None, total=None):
    payload = {
        "stage": stage,
        "percent": max(0, min(100, int(percent))),
        "message": message,
    }
    if current is not None:
        payload["current"] = float(current)
    if total is not None:
        payload["total"] = float(total)
    print(f"{PROGRESS_PREFIX}{json.dumps(payload, ensure_ascii=False)}", file=sys.stderr, flush=True)


def to_simplified(text):
    global _OPENCC_CONVERTER
    if not text:
        return text
    try:
        if _OPENCC_CONVERTER is None:
            from opencc import OpenCC
            _OPENCC_CONVERTER = OpenCC("t2s")
        return _OPENCC_CONVERTER.convert(text)
    except Exception:
        return text.translate(FALLBACK_T2S)


def find_ffmpeg():
    candidates = [
        os.environ.get("TRANSCRIBE_STUDIO_FFMPEG"),
        "/opt/homebrew/bin/ffmpeg",
        "/usr/local/bin/ffmpeg",
        shutil.which("ffmpeg"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    raise RuntimeError("未找到 ffmpeg。请先安装 Homebrew ffmpeg，或设置 TRANSCRIBE_STUDIO_FFMPEG。")


def find_ffprobe(ffmpeg):
    env_path = os.environ.get("TRANSCRIBE_STUDIO_FFPROBE")
    candidates = [
        env_path,
        str(Path(ffmpeg).with_name("ffprobe")) if ffmpeg else None,
        "/opt/homebrew/bin/ffprobe",
        "/usr/local/bin/ffprobe",
        shutil.which("ffprobe"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    raise RuntimeError("未找到 ffprobe。请确认 ffmpeg 安装完整。")


def extract_audio(ffmpeg, input_path, output_path):
    emit("正在抽取音频...")
    emit_progress("extract", 8, "正在抽取音频...")
    probe_cmd = [
        find_ffprobe(ffmpeg),
        "-v",
        "error",
        "-select_streams",
        "a",
        "-show_entries",
        "stream=index",
        "-of",
        "csv=p=0",
        str(input_path),
    ]
    probe = subprocess.run(probe_cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if not probe.stdout.strip():
        raise RuntimeError(
            "录制文件没有音频轨，无法转写。请在 macOS 的“屏幕与系统音频录制”和“麦克风”中授权本应用；"
            "如果仍无系统声音，请打开电脑外放，让应用使用麦克风兜底录音。"
        )
    cmd = [
        ffmpeg,
        "-y",
        "-i",
        str(input_path),
        "-vn",
        "-ac",
        "1",
        "-ar",
        "16000",
        "-f",
        "wav",
        str(output_path),
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    emit_progress("extract", 18, "音频抽取完成。")


def media_duration(ffprobe, input_path):
    cmd = [
        ffprobe,
        "-v",
        "error",
        "-show_entries",
        "format=duration",
        "-of",
        "default=noprint_wrappers=1:nokey=1",
        str(input_path),
    ]
    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        return max(0.0, float(result.stdout.strip()))
    except Exception:
        return 0.0


def has_video_stream(ffprobe, input_path):
    cmd = [
        ffprobe,
        "-v",
        "error",
        "-select_streams",
        "v",
        "-show_entries",
        "stream=index",
        "-of",
        "csv=p=0",
        str(input_path),
    ]
    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    return bool(result.stdout.strip())


def resolve_subtitle_ocr_bin():
    candidates = [
        os.environ.get("TRANSCRIBE_STUDIO_SUBTITLE_OCR"),
        str(Path(__file__).resolve().parents[1] / "native" / "bin" / "subtitle-ocr"),
        str(Path(sys.executable).resolve().parents[1] / "native" / "bin" / "subtitle-ocr"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    return ""


def extract_subtitle_frames(ffmpeg, ffprobe, input_path, output_dir, interval=1.2, max_frames=180):
    if not has_video_stream(ffprobe, input_path):
        return []
    duration = media_duration(ffprobe, input_path)
    if duration <= 0:
        return []
    frame_dir = output_dir / "subtitle_frames"
    frame_dir.mkdir(parents=True, exist_ok=True)
    count = min(max_frames, max(1, int(duration / interval)))
    timestamps = [min(duration - 0.1, index * interval + 0.2) for index in range(count)]
    frames = []
    emit("正在抽取画面字幕参考帧...")
    for index, timestamp in enumerate(timestamps):
        frame_path = frame_dir / f"subtitle-{index:04d}.png"
        cmd = [
            ffmpeg,
            "-y",
            "-ss",
            f"{timestamp:.2f}",
            "-i",
            str(input_path),
            "-frames:v",
            "1",
            "-vf",
            SUBTITLE_REGION_FILTER,
            str(frame_path),
        ]
        result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        if result.returncode == 0 and frame_path.exists():
            frames.append({"time": timestamp, "path": frame_path})
    return frames


def clean_subtitle_text(text):
    text = to_simplified(text)
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[|｜_~·•●]+", "", text)
    text = re.sub(r"^[0-9:：.,，。\\-— ]+", "", text)
    text = re.sub(r"[^\w\u4e00-\u9fff，。！？、：；“”‘’（）《》,.!?;:() -]+", "", text)
    return text.strip()


def similarity(a, b):
    if not a or not b:
        return 0.0
    aset = set(a)
    bset = set(b)
    return len(aset & bset) / max(1, len(aset | bset))


def normalized_for_dedupe(text):
    text = normalize_text(text)
    text = re.sub(r"[，。！？、：；,.!?;:\s]+", "", text)
    return text


def sequence_similarity(a, b):
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def dedupe_segments(segments, mode):
    if mode == "off":
        return segments, {"mode": mode, "removed": 0, "notes": []}
    window = 8 if mode == "normal" else 24
    threshold = 0.92 if mode == "normal" else 0.84
    min_chars = 8 if mode == "normal" else 6
    kept = []
    removed = []
    notes = []
    for segment in segments:
        current = normalized_for_dedupe(segment.get("text", ""))
        if len(current) < min_chars:
            kept.append(segment)
            continue
        duplicate_of = None
        for previous in reversed(kept[-window:]):
            prior = normalized_for_dedupe(previous.get("text", ""))
            if len(prior) < min_chars:
                continue
            same_text = current == prior
            close_text = sequence_similarity(current, prior) >= threshold
            overlap_text = len(current) >= 16 and (current in prior or prior in current)
            if same_text or close_text or overlap_text:
                duplicate_of = previous
                break
        if duplicate_of:
            removed.append(segment)
            if len(notes) < 12:
                notes.append(
                    f"[{int(segment['start'] // 60):02d}:{int(segment['start'] % 60):02d}] 疑似重复：{segment.get('text', '')[:80]}"
                )
            continue
        kept.append(segment)
    if removed:
        emit(f"重复内容去重：已移除 {len(removed)} 条疑似重复片段（模式：{mode}）。")
        emit_progress("dedupe", 78, f"重复内容去重完成：移除 {len(removed)} 条疑似重复片段。")
    return kept, {"mode": mode, "removed": len(removed), "notes": notes}


def recognize_subtitles(frames, output_dir):
    ocr_bin = resolve_subtitle_ocr_bin()
    raw_path = output_dir / "subtitle_ocr_raw.json"
    subtitle_path = output_dir / "字幕参考.txt"
    if not frames or not ocr_bin:
        raw_path.write_text("[]", encoding="utf-8")
        subtitle_path.write_text("", encoding="utf-8")
        return []

    emit("正在进行视频字幕 OCR 识别...")
    emit_progress("subtitle", 76, "正在进行视频字幕 OCR 识别...")
    results = []
    batch_size = 24
    for start in range(0, len(frames), batch_size):
        batch = frames[start:start + batch_size]
        cmd = [ocr_bin, *[str(item["path"]) for item in batch]]
        result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=180)
        if result.returncode != 0:
            emit(f"字幕 OCR 批次失败：{result.stderr.strip() or result.stdout.strip()}")
            continue
        try:
            payload = json.loads(result.stdout or "[]")
        except json.JSONDecodeError:
            payload = []
        by_path = {item.get("path"): item.get("text", "") for item in payload if isinstance(item, dict)}
        for item in batch:
            text = clean_subtitle_text(by_path.get(str(item["path"]), ""))
            if len(text) >= 2:
                results.append({"time": item["time"], "text": text})

    merged = []
    for item in results:
        if merged and (item["text"] == merged[-1]["text"] or similarity(item["text"], merged[-1]["text"]) > 0.82):
            continue
        merged.append(item)

    raw_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    subtitle_lines = [
        f"[{int(item['time'] // 60):02d}:{int(item['time'] % 60):02d}] {item['text']}"
        for item in merged
    ]
    subtitle_path.write_text("\n".join(subtitle_lines), encoding="utf-8")
    emit_progress("subtitle", 84, f"字幕辅助识别完成：{len(merged)} 条参考字幕。")
    return merged


def load_engine(model_name, language):
    try:
        from faster_whisper import WhisperModel

        def run(audio_path, duration=0):
            emit(f"[语音转文字] 使用 faster-whisper 模型：{model_name}")
            emit_progress("transcribe", 24, "正在加载语音识别模型...")
            model = WhisperModel(model_name, device="auto", compute_type="auto")
            kwargs = {"vad_filter": True}
            if language and language != "auto":
                kwargs["language"] = language
            segments, info = model.transcribe(str(audio_path), **kwargs)
            results = []
            last_percent = 0
            for segment in segments:
                item = {
                    "start": float(segment.start),
                    "end": float(segment.end),
                    "text": segment.text.strip(),
                }
                results.append(item)
                if duration > 0:
                    percent = 25 + min(48, int((item["end"] / duration) * 48))
                    if percent != last_percent:
                        last_percent = percent
                        emit_progress(
                            "transcribe",
                            percent,
                            "正在进行 AI 语音识别...",
                            current=min(item["end"], duration),
                            total=duration,
                        )
            emit_progress("transcribe", 74, "AI 语音识别完成。", current=duration, total=duration if duration > 0 else None)
            return results, getattr(info, "language", language or "auto")

        return run
    except Exception as first_error:
        try:
            import whisper

            def run(audio_path, duration=0):
                emit(f"[语音转文字] 使用 openai-whisper 模型：{model_name}")
                emit_progress("transcribe", 24, "正在加载语音识别模型...")
                model = whisper.load_model(model_name)
                kwargs = {}
                if language and language != "auto":
                    kwargs["language"] = language
                result = model.transcribe(str(audio_path), **kwargs)
                segments = []
                for segment in result.get("segments", []):
                    item = {
                        "start": float(segment.get("start", 0)),
                        "end": float(segment.get("end", 0)),
                        "text": segment.get("text", "").strip(),
                    }
                    segments.append(item)
                    if duration > 0:
                        emit_progress(
                            "transcribe",
                            25 + min(48, int((item["end"] / duration) * 48)),
                            "正在进行 AI 语音识别...",
                            current=min(item["end"], duration),
                            total=duration,
                        )
                emit_progress("transcribe", 74, "AI 语音识别完成。", current=duration, total=duration if duration > 0 else None)
                return segments, result.get("language", language or "auto")

            return run
        except Exception as second_error:
            raise RuntimeError(
                "未检测到本机 Whisper AI 转写引擎。可执行：\n"
                "python3 -m pip install faster-whisper python-docx\n"
                "或：python3 -m pip install openai-whisper python-docx\n"
                f"faster-whisper 错误：{first_error}\nopenai-whisper 错误：{second_error}"
            )


def normalize_text(text):
    text = to_simplified(text)
    text = re.sub(r"\s+", " ", text).strip()
    for pattern in FILLER_PATTERNS:
        text = re.sub(pattern, "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"([。！？!?])\1+", r"\1", text)
    return to_simplified(text)


def clean_segments(segments, style):
    cleaned = []
    for segment in segments:
        text = normalize_text(segment["text"])
        if not text or len(text) <= 1:
            continue
        if style != "verbatim" and re.fullmatch(r"[啊嗯呃额哈\s，。,.!?！？]+", text):
            continue
        cleaned.append({**segment, "text": text})
    return cleaned


def split_sentences(text):
    parts = re.split(r"(?<=[。！？!?])\s*", text)
    return [part.strip() for part in parts if part.strip()]


def template_label(document_template):
    return DOCUMENT_TEMPLATES.get(document_template, DOCUMENT_TEMPLATES["general"])["label"]


def build_document(cleaned, style, document_template="general"):
    full_text = "\n".join(segment["text"] for segment in cleaned).strip()
    sentences = split_sentences(full_text)
    if style == "outline":
        important = sentences[:80]
        body = "\n".join(f"- {sentence}" for sentence in important)
    elif style == "verbatim":
        body = full_text
    else:
        important = []
        for sentence in sentences:
            if len(sentence) < 6:
                continue
            if sentence not in important:
                important.append(sentence)
        body = "\n".join(important)
    if document_template != "general" and body:
        label = template_label(document_template)
        body = f"## {label}\n\n{body}"
    return body.strip() or full_text


def build_organizer_prompt(body, timeline, style, subtitle_reference="", dedupe_report=None, document_template="general"):
    style_hint = {
        "clean": "清理废话、口头禅、重复句和笑声，保留重点信息，输出适合阅读的简体中文整理稿。",
        "outline": "提炼成有层次的简体中文提纲，保留关键观点、结论和行动项。",
        "verbatim": "尽量保留原意和顺序，只修正识别错误、繁体字、废话和明显重复。"
    }.get(style, "整理成简体中文文档。")
    template = DOCUMENT_TEMPLATES.get(document_template, DOCUMENT_TEMPLATES["general"])
    source = body or timeline
    source = source[:12000]
    subtitle_block = subtitle_reference[:6000] if subtitle_reference else "无"
    dedupe_hint = "未开启。"
    if dedupe_report:
        removed = int(dedupe_report.get("removed") or 0)
        mode = dedupe_report.get("mode") or "off"
        notes = "\n".join(dedupe_report.get("notes") or [])
        dedupe_hint = f"去重模式：{mode}；已预处理移除疑似重复片段：{removed} 条。"
        if notes:
            dedupe_hint = f"{dedupe_hint}\n疑似重复片段示例：\n{notes}"
    return textwrap.dedent(f"""
    你是“零创AI 智能转写器”的文档整理助手。
    请基于下面的语音转写内容进行二次整理，并参考视频画面字幕。
    要求：
    1. 全部输出简体中文。
    2. {style_hint}
    3. 文档模板：{template["label"]}。{template["hint"]}
    4. 删除无意义语气词、笑声、掌声、背景音乐提示和明显废话。
    5. 不要编造原文没有的信息。
    6. 如果语音识别和画面字幕冲突，优先根据上下文判断；专有名词、人名、课程术语可优先参考字幕。
    7. 如果发现在线播放卡顿、回放、跳回开头导致的重复段落，只保留一次；但对主播有意强调的重点不要过度删除。
    8. 如果发现上下文明显断裂、缺少承接，保留可确认内容，并在末尾用一句话标注“可能存在因播放卡顿造成的内容缺失”。
    9. 只输出最终整理稿，不要解释你的处理过程。

    重复/缺失处理参考：
    {dedupe_hint}

    语音转写内容：
    {source}

    画面字幕参考：
    {subtitle_block}
    """).strip()


def run_openclaw_organizer(body, timeline, style, openclaw_bin, openclaw_model, subtitle_reference="", dedupe_report=None, document_template="general"):
    prompt = build_organizer_prompt(body, timeline, style, subtitle_reference, dedupe_report, document_template)
    cmd = [
        openclaw_bin,
        "infer",
        "model",
        "run",
        "--gateway",
        "--json",
        "--prompt",
        prompt,
    ]
    if openclaw_model:
        cmd.extend(["--model", openclaw_model])
    emit(f"[智能整理] 正在调用 OpenClaw/Qwen 增强整理（模型：{openclaw_model or '默认'}）...")
    emit_progress("organize", 86, "正在调用 OpenClaw/Qwen 增强整理...")
    env = os.environ.copy()
    extra_path = "/opt/homebrew/bin:/opt/homebrew/sbin:/usr/local/bin:/usr/local/sbin:/usr/bin:/bin:/usr/sbin:/sbin"
    env["PATH"] = f"{env.get('PATH', '')}:{extra_path}" if env.get("PATH") else extra_path
    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=600, env=env)
    if result.returncode != 0:
        raise RuntimeError((result.stderr or result.stdout or "OpenClaw 整理失败").strip())
    text = parse_openclaw_output(result.stdout)
    text = to_simplified(text).strip()
    if not text:
        raise RuntimeError("OpenClaw 没有返回有效整理内容。")
    return text


def run_local_ai_organizer(body, timeline, style, local_ai_base_url, local_ai_model, subtitle_reference="", dedupe_report=None, document_template="general"):
    if not local_ai_base_url:
        raise RuntimeError("本地大模型地址为空。")
    if not local_ai_model:
        raise RuntimeError("本地大模型名称为空。")
    prompt = build_organizer_prompt(body, timeline, style, subtitle_reference, dedupe_report, document_template)
    base_url = local_ai_base_url.rstrip("/")
    if base_url.endswith("/v1"):
        endpoint = f"{base_url}/chat/completions"
    else:
        endpoint = f"{base_url}/v1/chat/completions"
    payload = {
        "model": local_ai_model,
        "messages": [
            {"role": "system", "content": "你是严谨的中文转写文档整理助手，只输出最终整理稿。"},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.2,
        "stream": False
    }
    emit(f"[智能整理] 正在调用本地大模型直连整理（模型：{local_ai_model}）...")
    attempts = [600, 900, 1200]
    errors = []
    for index, timeout in enumerate(attempts, start=1):
        emit(f"本地大模型第 {index}/3 次请求，最长等待 {timeout // 60} 分钟...")
        emit_progress("organize", 84 + index, f"本地大模型第 {index}/3 次请求：{local_ai_model}")
        request = urllib.request.Request(
            endpoint,
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST"
        )
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                data = response.read().decode("utf-8", errors="replace")
            response_payload = json.loads(data or "{}")
            text = to_simplified(extract_text(response_payload)).strip()
            if not text:
                raise RuntimeError("本地大模型没有返回有效整理内容。")
            return text
        except urllib.error.HTTPError as error:
            detail = error.read().decode("utf-8", errors="replace") if error.fp else ""
            errors.append(f"第 {index} 次：HTTP {error.code} {detail[:800]}")
        except (urllib.error.URLError, socket.timeout, TimeoutError, json.JSONDecodeError, RuntimeError) as error:
            errors.append(f"第 {index} 次：{error}")
    raise RuntimeError("本地大模型连续 3 次未返回有效整理内容；" + "；".join(errors))


def chat_completion_endpoint(base_url):
    value = (base_url or "").rstrip("/")
    if not value:
        return ""
    if value.endswith("/v1"):
        return f"{value}/chat/completions"
    return f"{value}/v1/chat/completions"


def run_cloud_ai_organizer(body, timeline, style, cloud_ai_base_url, cloud_ai_model, subtitle_reference="", dedupe_report=None, document_template="general"):
    api_key = os.environ.get("TRANSCRIBE_STUDIO_CLOUD_AI_API_KEY", "").strip()
    if not cloud_ai_base_url:
        raise RuntimeError("云端 API 地址为空。")
    if not cloud_ai_model:
        raise RuntimeError("云端模型名称为空。")
    if not api_key:
        raise RuntimeError("云端 API Key 为空。")
    prompt = build_organizer_prompt(body, timeline, style, subtitle_reference, dedupe_report, document_template)
    endpoint = chat_completion_endpoint(cloud_ai_base_url)
    payload = {
        "model": cloud_ai_model,
        "messages": [
            {"role": "system", "content": "你是严谨的中文转写文档整理助手，只输出最终整理稿。"},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.2,
        "stream": False
    }
    emit(f"[智能整理] 正在调用云端大模型 API 整理（模型：{cloud_ai_model}）...")
    emit_progress("organize", 86, f"正在调用云端大模型整理：{cloud_ai_model}")
    request = urllib.request.Request(
        endpoint,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}",
        },
        method="POST"
    )
    try:
        with urllib.request.urlopen(request, timeout=600) as response:
            data = response.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace") if error.fp else ""
        raise RuntimeError(f"云端大模型请求失败：HTTP {error.code} {detail[:800]}") from error
    payload = json.loads(data or "{}")
    text = to_simplified(extract_text(payload)).strip()
    if not text:
        raise RuntimeError("云端大模型没有返回有效整理内容。")
    return text


def parse_openclaw_output(stdout):
    stripped = stdout.strip()
    if not stripped:
        return ""
    try:
        payload = json.loads(stripped)
        return extract_text(payload)
    except json.JSONDecodeError:
        pass
    for line in reversed([line.strip() for line in stripped.splitlines() if line.strip()]):
        try:
            payload = json.loads(line)
            text = extract_text(payload)
            if text:
                return text
        except json.JSONDecodeError:
            continue
    return stripped


def extract_text(value):
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        for key in ("text", "outputs", "output", "content", "message", "reply", "result", "data"):
            if key in value:
                text = extract_text(value[key])
                if text:
                    return text
        choices = value.get("choices")
        if isinstance(choices, list):
            return extract_text(choices)
        if isinstance(value.get("message"), dict):
            return extract_text(value["message"])
    if isinstance(value, list):
        parts = [extract_text(item) for item in value]
        return "\n".join(part for part in parts if part)
    return ""


def organize_body(body, timeline, style, organizer, openclaw_bin, openclaw_model, local_ai_base_url="", local_ai_model="", cloud_ai_base_url="", cloud_ai_model="", subtitle_reference="", dedupe_report=None, document_template="general"):
    report = {
        "organizer_requested": organizer,
        "organizer_actual": "local",
        "style": style,
        "document_template": document_template,
        "openclaw_bin": openclaw_bin if organizer == "openclaw" else "",
        "openclaw_model": openclaw_model if organizer == "openclaw" else "",
        "local_ai_base_url": local_ai_base_url if organizer == "localai" else "",
        "local_ai_model": local_ai_model if organizer == "localai" else "",
        "cloud_ai_base_url": cloud_ai_base_url if organizer == "cloudai" else "",
        "cloud_ai_model": cloud_ai_model if organizer == "cloudai" else "",
        "failure_reason": "",
    }
    if organizer not in ("openclaw", "localai", "cloudai"):
        emit_progress("organize", 88, "正在进行本机智能整理...")
        return body, "local", report
    if organizer == "openclaw":
        try:
            enhanced = run_openclaw_organizer(body, timeline, style, openclaw_bin, openclaw_model, subtitle_reference, dedupe_report, document_template)
            emit_progress("organize", 92, "OpenClaw/Qwen 增强整理完成。")
            report["organizer_actual"] = "openclaw"
            return enhanced, "openclaw", report
        except Exception as error:
            emit(f"OpenClaw 增强整理不可用，已保留本机整理结果：{error}")
            emit_progress("organize", 90, "OpenClaw 不可用，已回退本机整理。")
            report["organizer_actual"] = "local-fallback"
            report["failure_reason"] = str(error)
            return body, "local-fallback", report
    if organizer == "cloudai":
        try:
            enhanced = run_cloud_ai_organizer(body, timeline, style, cloud_ai_base_url, cloud_ai_model, subtitle_reference, dedupe_report, document_template)
            emit_progress("organize", 92, "云端大模型增强整理完成。")
            report["organizer_actual"] = "cloudai"
            return enhanced, "cloudai", report
        except Exception as error:
            emit(f"云端大模型增强整理不可用，已保留本机整理结果：{error}")
            emit_progress("organize", 90, "云端大模型不可用，已回退本机整理。")
            report["organizer_actual"] = "cloudai-fallback"
            report["failure_reason"] = str(error)
            return body, "cloudai-fallback", report
    try:
        enhanced = run_local_ai_organizer(body, timeline, style, local_ai_base_url, local_ai_model, subtitle_reference, dedupe_report, document_template)
        emit_progress("organize", 92, "本地大模型增强整理完成。")
        report["organizer_actual"] = "localai"
        return enhanced, "localai", report
    except Exception as error:
        emit(f"本地大模型增强整理不可用，已保留本机整理结果：{error}")
        emit_progress("organize", 90, "本地大模型不可用，已回退本机整理。")
        report["organizer_actual"] = "localai-fallback"
        report["failure_reason"] = str(error)
        return body, "localai-fallback", report


def report_label(value):
    labels = {
        "local": "本机规则整理",
        "openclaw": "OpenClaw/Qwen 增强整理",
        "localai": "本地大模型直连",
        "cloudai": "云端大模型 API",
        "local-fallback": "OpenClaw 失败后回退本机规则整理",
        "localai-fallback": "本地大模型失败后回退本机规则整理",
        "cloudai-fallback": "云端大模型失败后回退本机规则整理",
    }
    return labels.get(value, value or "未记录")


def build_processing_report_lines(language, style, organizer_report, dedupe_report=None, subtitle_count=0):
    lines = [
        f"识别语言：{language}",
        f"整理方式：{style}",
        f"文档模板：{template_label(organizer_report.get('document_template') or 'general')}",
        f"智能整理请求：{report_label(organizer_report.get('organizer_requested'))}",
        f"智能整理实际：{report_label(organizer_report.get('organizer_actual'))}",
    ]
    if organizer_report.get("local_ai_base_url"):
        lines.append(f"本地模型地址：{organizer_report.get('local_ai_base_url')}")
    if organizer_report.get("local_ai_model"):
        lines.append(f"本地模型名称：{organizer_report.get('local_ai_model')}")
    if organizer_report.get("openclaw_bin"):
        lines.append(f"OpenClaw 命令：{organizer_report.get('openclaw_bin')}")
    if organizer_report.get("openclaw_model"):
        lines.append(f"OpenClaw 模型：{organizer_report.get('openclaw_model')}")
    if organizer_report.get("cloud_ai_base_url"):
        lines.append(f"云端 API 地址：{organizer_report.get('cloud_ai_base_url')}")
    if organizer_report.get("cloud_ai_model"):
        lines.append(f"云端模型名称：{organizer_report.get('cloud_ai_model')}")
    if organizer_report.get("failure_reason"):
        lines.append(f"失败原因：{organizer_report.get('failure_reason')}")
    else:
        lines.append("失败原因：无")
    if dedupe_report:
        lines.append(f"重复内容去重：{dedupe_report.get('mode')}；移除 {dedupe_report.get('removed', 0)} 条")
    lines.append(f"字幕辅助识别数量：{subtitle_count} 条")
    return [to_simplified(line) for line in lines]


def write_outputs(output_dir, source_path, cleaned, language, style, organizer, openclaw_bin, openclaw_model, local_ai_base_url="", local_ai_model="", cloud_ai_base_url="", cloud_ai_model="", subtitles=None, dedupe_report=None, document_template="general"):
    emit_progress("document", 90, "正在生成转写文档...")
    title = to_simplified(source_path.stem)
    cleaned = [{**item, "text": to_simplified(item["text"])} for item in cleaned]
    timeline = "\n".join(
        f"[{int(item['start'] // 60):02d}:{int(item['start'] % 60):02d}] {item['text']}"
        for item in cleaned
    )
    subtitle_reference = "\n".join(
        f"[{int(item['time'] // 60):02d}:{int(item['time'] % 60):02d}] {item['text']}"
        for item in (subtitles or [])
    )
    body = to_simplified(build_document(cleaned, style, document_template))
    if subtitle_reference and organizer not in ("openclaw", "localai", "cloudai"):
        body = f"{body}\n\n## 字幕参考校正\n\n{subtitle_reference}".strip()
    body, organizer_used, organizer_report = organize_body(
        body,
        timeline,
        style,
        organizer,
        openclaw_bin,
        openclaw_model,
        local_ai_base_url,
        local_ai_model,
        cloud_ai_base_url,
        cloud_ai_model,
        subtitle_reference,
        dedupe_report,
        document_template
    )
    body = to_simplified(body)
    dedupe_summary = ""
    if dedupe_report and dedupe_report.get("mode") != "off":
        notes = "\n".join(dedupe_report.get("notes") or [])
        dedupe_summary = (
            f"模式：{dedupe_report.get('mode')}；已移除疑似重复片段：{dedupe_report.get('removed', 0)} 条。"
            + (f"\n\n{notes}" if notes else "")
        )
    report_lines = build_processing_report_lines(
        language,
        style,
        organizer_report,
        dedupe_report,
        len(subtitles or [])
    )
    md_report = "\n".join(f"- {line}" for line in report_lines)
    txt_report = "\n".join(report_lines)

    markdown = output_dir / "转写文档.md"
    txt = output_dir / "转写文档.txt"
    docx = output_dir / "转写文档.docx"

    md_content = (
        f"# {title}\n\n"
        f"## 整理稿\n\n{body}\n\n"
        f"## 处理报告\n\n{md_report}\n\n"
        f"## 重复内容处理\n\n{dedupe_summary or '未开启重复内容去重。'}\n\n"
        f"## 画面字幕参考\n\n{subtitle_reference or '未识别到可用字幕。'}\n\n"
        f"## 时间线原文\n\n{timeline}\n"
    )
    markdown.write_text(md_content, encoding="utf-8")
    txt.write_text(
        f"{title}\n\n整理稿\n\n{body}\n\n处理报告\n\n{txt_report}\n\n重复内容处理\n\n{dedupe_summary or '未开启重复内容去重。'}\n\n画面字幕参考\n\n{subtitle_reference or '未识别到可用字幕。'}\n\n时间线原文\n\n{timeline}\n",
        encoding="utf-8"
    )

    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_heading("整理稿", level=2)
    for paragraph in body.splitlines():
        if paragraph.startswith("- "):
            doc.add_paragraph(paragraph[2:], style="List Bullet")
        elif paragraph.strip():
            doc.add_paragraph(paragraph)
    doc.add_heading("处理报告", level=2)
    for line in report_lines:
        doc.add_paragraph(line)
    if subtitle_reference:
        doc.add_heading("画面字幕参考", level=2)
        for line in subtitle_reference.splitlines():
            doc.add_paragraph(line)
    doc.add_heading("重复内容处理", level=2)
    doc.add_paragraph(dedupe_summary or "未开启重复内容去重。")
    doc.add_heading("时间线原文", level=2)
    for line in timeline.splitlines():
        doc.add_paragraph(line)
    doc.add_paragraph(f"识别语言：{language}；整理方式：{style}；智能整理：{organizer_used}")
    doc.save(docx)
    emit_progress("document", 98, "转写文档已生成。")

    return markdown, txt, docx, body, organizer_used, organizer_report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--language", default="zh")
    parser.add_argument("--model", default="small")
    parser.add_argument("--style", default="clean")
    parser.add_argument("--organizer", default="local", choices=["local", "openclaw", "localai", "cloudai"])
    parser.add_argument("--openclaw-bin", default="openclaw")
    parser.add_argument("--openclaw-model", default="")
    parser.add_argument("--local-ai-base-url", default="")
    parser.add_argument("--local-ai-model", default="")
    parser.add_argument("--cloud-ai-base-url", default="")
    parser.add_argument("--cloud-ai-model", default="")
    parser.add_argument("--subtitle-mode", default="off", choices=["off", "auto"])
    parser.add_argument("--dedupe-mode", default="normal", choices=["off", "normal", "strong"])
    parser.add_argument("--document-template", default="general", choices=list(DOCUMENT_TEMPLATES.keys()))
    parser.add_argument("--segments-json", default="")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    if args.segments_json:
        emit("正在复用已有逐字稿换模板生成文档...")
        emit_progress("organize", 72, "正在复用已有逐字稿换模板生成文档...")
        segments_path = Path(args.segments_json).expanduser().resolve()
        cleaned = json.loads(segments_path.read_text(encoding="utf-8"))
        cleaned = clean_segments(cleaned, args.style)
        cleaned = [{**item, "text": to_simplified(item["text"])} for item in cleaned]
        detected_language = args.language or "zh"
        dedupe_report = {"mode": "off", "removed": 0, "notes": ["换模板生成复用已处理逐字稿，未重新执行语音识别和去重。"]}
        subtitles = []
        subtitle_reference_file = output_dir / "字幕参考.txt"
        subtitle_reference_file.write_text("", encoding="utf-8")
    else:
        ffmpeg = find_ffmpeg()
        ffprobe = find_ffprobe(ffmpeg)
        duration = media_duration(ffprobe, input_path)
        audio_path = output_dir / "audio.wav"
        extract_audio(ffmpeg, input_path, audio_path)

        engine = load_engine(args.model, args.language)
        emit("正在进行 AI 语音识别...")
        emit_progress("transcribe", 22, "正在进行 AI 语音识别...", current=0, total=duration if duration > 0 else None)
        segments, detected_language = engine(audio_path, duration)
        cleaned = clean_segments(segments, args.style)
        cleaned = [{**item, "text": to_simplified(item["text"])} for item in cleaned]
        cleaned, dedupe_report = dedupe_segments(cleaned, args.dedupe_mode)

        subtitles = []
        subtitle_reference_file = output_dir / "字幕参考.txt"
        if args.subtitle_mode == "auto":
            try:
                frames = extract_subtitle_frames(ffmpeg, ffprobe, input_path, output_dir)
                subtitles = recognize_subtitles(frames, output_dir)
                emit(f"字幕辅助识别完成：{len(subtitles)} 条参考字幕。")
            except Exception as error:
                subtitle_reference_file.write_text("", encoding="utf-8")
                emit(f"字幕辅助识别不可用，已继续使用纯音频转写：{error}")
        else:
            subtitle_reference_file.write_text("", encoding="utf-8")

    raw_json = output_dir / "segments.json"
    raw_json.write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")

    markdown, txt, docx, body, organizer_used, organizer_report = write_outputs(
        output_dir,
        input_path,
        cleaned,
        detected_language,
        args.style,
        args.organizer,
        args.openclaw_bin,
        args.openclaw_model,
        args.local_ai_base_url,
        args.local_ai_model,
        args.cloud_ai_base_url,
        args.cloud_ai_model,
        subtitles,
        dedupe_report,
        args.document_template,
    )
    summary = body.splitlines()[0][:120] if body else "已完成转写。"
    emit_progress("done", 100, "转写完成。")
    print(json.dumps({
        "output_dir": str(output_dir),
        "markdown": str(markdown),
        "txt": str(txt),
        "docx": str(docx),
        "segments": str(raw_json),
        "summary": summary,
        "language": detected_language,
        "organizer": organizer_used,
        "processing_report": organizer_report,
        "dedupe": dedupe_report,
        "subtitles": str(subtitle_reference_file),
        "subtitle_count": len(subtitles),
    }, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    try:
        main()
    except subprocess.CalledProcessError as error:
        raise SystemExit(f"音频处理失败：{error.stderr.decode('utf-8', errors='ignore')}")
    except Exception as error:
        raise SystemExit(str(error))
